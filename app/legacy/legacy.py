from __future__ import annotations

import calendar
import hashlib
import io
import json
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.cell import range_boundaries  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from app.legacy.types import (
    DatePrecision,
    FieldEvidence,
    LegacyExtraction,
    LegacyRecord,
    Scalar,
    SourceLocation,
)

LEGACY_SOURCE_PATHS = {
    "income": Path("收入表格/收入明细.xlsx"),
    "account_purchase": Path("支出表格/账号购买明细.xlsx"),
    "subscription": Path("支出表格/账号订阅明细.xlsx"),
    "other_expense": Path("支出表格/其他费用明细.xlsx"),
    "cloud_service": Path("支出表格/云服务运营明细.xlsx"),
    "upstream": Path("支出表格/上游中转站充值明细.xlsx"),
}

_MONTH_PATTERN = re.compile(r"^(\d{4})-(\d{2})$")
_DATE_PATTERN = re.compile(r"^(\d{4}-\d{2}-\d{2})$")
_DATETIME_PATTERN = re.compile(r"^(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}(?::\d{2})?)$")
_CENT = Decimal("0.01")


def _decimal_text(value: int | float | Decimal) -> str:
    text = format(Decimal(str(value)), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def normalize_scalar(value: Any) -> Scalar:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        return _decimal_text(value)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def read_source_bytes(path: str | Path) -> bytes:
    source = Path(path)
    try:
        return source.read_bytes()
    except PermissionError:
        if os.name != "nt":
            raise
        return _read_windows_shared(source)


def _read_windows_shared(path: Path) -> bytes:
    import ctypes
    import msvcrt

    ctypes_windows: Any = ctypes
    msvcrt_windows: Any = msvcrt

    kernel32 = ctypes_windows.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        ctypes.c_wchar_p,
        ctypes.c_uint32,
        ctypes.c_uint32,
        ctypes.c_void_p,
        ctypes.c_uint32,
        ctypes.c_uint32,
        ctypes.c_void_p,
    ]
    create_file.restype = ctypes.c_void_p
    close_handle = kernel32.CloseHandle
    close_handle.argtypes = [ctypes.c_void_p]
    close_handle.restype = ctypes.c_int

    handle = create_file(
        str(path),
        0x80000000,
        0x00000001 | 0x00000002 | 0x00000004,
        None,
        3,
        0x00000080,
        None,
    )
    if handle == ctypes.c_void_p(-1).value:
        error_code = ctypes_windows.get_last_error()
        raise PermissionError(error_code, os.strerror(error_code), str(path))

    try:
        descriptor = msvcrt_windows.open_osfhandle(handle, os.O_RDONLY)
    except Exception:
        close_handle(handle)
        raise
    with os.fdopen(descriptor, "rb") as stream:
        return stream.read()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@dataclass(frozen=True, slots=True)
class _WorkbookContext:
    source_label: str
    sha256: str
    formulas: Workbook
    values: Workbook


def extract_legacy_workbooks(workspace_root: str | Path) -> LegacyExtraction:
    root = Path(workspace_root).resolve()
    contexts = {
        name: _load_context(root, relative_path)
        for name, relative_path in LEGACY_SOURCE_PATHS.items()
    }

    records: list[LegacyRecord] = []
    records.extend(_extract_income(contexts["income"]))
    records.extend(_extract_account_purchases(contexts["account_purchase"]))
    records.extend(_extract_subscriptions(contexts["subscription"]))
    records.extend(_extract_other_expenses(contexts["other_expense"]))
    records.extend(_extract_cloud_services(contexts["cloud_service"]))
    records.extend(_extract_upstream_recharges(contexts["upstream"]))

    fingerprints = [record.fingerprint for record in records]
    if len(fingerprints) != len(set(fingerprints)):
        raise ValueError("Legacy extraction produced duplicate record fingerprints")

    return LegacyExtraction(
        records=tuple(records),
        source_hashes={context.source_label: context.sha256 for context in contexts.values()},
    )


def _load_context(root: Path, relative_path: Path) -> _WorkbookContext:
    path = root / relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Legacy workbook not found: {path}")
    data = read_source_bytes(path)
    return _WorkbookContext(
        source_label=relative_path.as_posix(),
        sha256=sha256_bytes(data),
        formulas=load_workbook(
            io.BytesIO(data), data_only=False, read_only=False, keep_links=False
        ),
        values=load_workbook(io.BytesIO(data), data_only=True, read_only=False, keep_links=False),
    )


def _extract_income(context: _WorkbookContext) -> list[LegacyRecord]:
    configurations = (
        ("Stripe收入", "StripeIncomeTable", "stripe", "payment"),
        ("xorpay收入", "XorpayIncomeTable", "xorpay", "daily_summary"),
        ("链动提现", "LiandongPayoutTable", "liandong", "payout"),
    )
    records: list[LegacyRecord] = []
    for sheet_name, table_name, source_system, recognition_method in configurations:
        formula_sheet = context.formulas[sheet_name]
        value_sheet = context.values[sheet_name]
        for row_number in _business_rows(formula_sheet, value_sheet, table_name, (1, 4)):
            original_amount = _required_decimal(value_sheet.cell(row_number, 4).value)
            currency = _required_text(value_sheet.cell(row_number, 5).value).upper()
            original_fee = _optional_decimal(value_sheet.cell(row_number, 6).value)
            fee_currency = _optional_text(value_sheet.cell(row_number, 7).value) or currency
            fee_currency = fee_currency.upper()

            gross = _optional_decimal(value_sheet.cell(row_number, 8).value)
            if gross is None:
                gross = _round_money(original_amount * _income_fx_rate(value_sheet, currency))
            fee = _optional_decimal(value_sheet.cell(row_number, 9).value)
            if fee is None:
                if original_fee is None and source_system == "xorpay":
                    original_fee = _round_money(
                        original_amount * _required_decimal(value_sheet["D3"].value)
                    )
                    fee_currency = "CNY"
                fee = _round_money(
                    (original_fee or Decimal("0")) * _income_fx_rate(value_sheet, fee_currency)
                )
            net = _optional_decimal(value_sheet.cell(row_number, 10).value)
            if net is None:
                net = _round_money(gross - fee)

            date_fields = _normalize_date(value_sheet.cell(row_number, 1).value)
            external_id = _optional_text(value_sheet.cell(row_number, 2).value)
            status = _optional_text(value_sheet.cell(row_number, 3).value)
            notes = _optional_text(value_sheet.cell(row_number, 11).value)
            fx_rate = _income_fx_rate(value_sheet, currency)
            fee_fx_rate = _income_fx_rate(value_sheet, fee_currency)
            business = {
                "record_type": "income",
                "source_system": source_system,
                "accounting_date": date_fields[0],
                "occurred_at": date_fields[1],
                "external_id": external_id,
                "gross": _decimal_string(gross),
                "fee": _decimal_string(fee),
                "net": _decimal_string(net),
            }
            source_fp, business_fp, fingerprint = _fingerprints(
                context=context,
                sheet=sheet_name,
                row=row_number,
                cell=None,
                business=business,
                trusted_external_id=external_id,
            )
            records.append(
                LegacyRecord(
                    record_type="income",
                    source_system=source_system,
                    category="channel_income",
                    accounting_date=date_fields[0],
                    occurred_at=date_fields[1],
                    date_precision=date_fields[2],
                    service_period=date_fields[3],
                    description=f"{sheet_name}收入",
                    counterparty=None,
                    channel=source_system,
                    external_id=external_id,
                    status=status,
                    currency=currency,
                    original_amount=_decimal_string(original_amount),
                    fx_rate=_decimal_string(fx_rate),
                    gross_amount_cny=_decimal_string(gross),
                    fee_amount_cny=_decimal_string(fee),
                    net_amount_cny=_decimal_string(net),
                    operating_amount_cny=_decimal_string(gross),
                    cash_amount_cny=_decimal_string(net),
                    credited_value=None,
                    recognition_method=recognition_method,
                    notes=notes,
                    fingerprint=fingerprint,
                    business_fingerprint=business_fp,
                    source_fingerprint=source_fp,
                    evidence=_row_evidence(
                        context,
                        sheet_name,
                        row_number,
                        {
                            "occurred_at": 1,
                            "external_id": 2,
                            "status": 3,
                            "original_amount": 4,
                            "currency": 5,
                            "original_fee": 6,
                            "fee_currency": 7,
                            "gross_amount_cny": 8,
                            "fee_amount_cny": 9,
                            "net_amount_cny": 10,
                            "notes": 11,
                        },
                    ),
                    metadata={
                        "original_fee": _decimal_string(original_fee or Decimal("0")),
                        "fee_currency": fee_currency,
                        "fee_fx_rate": _decimal_string(fee_fx_rate),
                    },
                )
            )
    return records


def _extract_account_purchases(context: _WorkbookContext) -> list[LegacyRecord]:
    sheet_name = "账号购买"
    formula_sheet = context.formulas[sheet_name]
    value_sheet = context.values[sheet_name]
    records: list[LegacyRecord] = []
    for row_number in _business_rows(formula_sheet, value_sheet, "AccountPurchaseTable", (4, 5)):
        amount = _required_decimal(value_sheet.cell(row_number, 5).value)
        date_fields = _normalize_date(value_sheet.cell(row_number, 1).value)
        channel = _optional_text(value_sheet.cell(row_number, 2).value)
        description = _required_text(value_sheet.cell(row_number, 3).value)
        external_id = _required_text(value_sheet.cell(row_number, 4).value)
        status = _optional_text(value_sheet.cell(row_number, 6).value)
        business = {
            "record_type": "expense",
            "source_system": "account_purchase",
            "external_id": external_id,
            "amount": _decimal_string(amount),
        }
        source_fp, business_fp, fingerprint = _fingerprints(
            context=context,
            sheet=sheet_name,
            row=row_number,
            cell=None,
            business=business,
            trusted_external_id=external_id,
        )
        records.append(
            LegacyRecord(
                record_type="expense",
                source_system="account_purchase",
                category="account_purchase",
                accounting_date=date_fields[0],
                occurred_at=date_fields[1],
                date_precision=date_fields[2],
                service_period=date_fields[3],
                description=description,
                counterparty=channel,
                channel=channel,
                external_id=external_id,
                status=status,
                currency="CNY",
                original_amount=_decimal_string(amount),
                fx_rate="1",
                gross_amount_cny=None,
                fee_amount_cny=None,
                net_amount_cny=None,
                operating_amount_cny=_decimal_string(amount),
                cash_amount_cny=_decimal_string(amount),
                credited_value=None,
                recognition_method=None,
                notes=None,
                fingerprint=fingerprint,
                business_fingerprint=business_fp,
                source_fingerprint=source_fp,
                evidence=_row_evidence(
                    context,
                    sheet_name,
                    row_number,
                    {
                        "occurred_at": 1,
                        "channel": 2,
                        "description": 3,
                        "external_id": 4,
                        "amount_cny": 5,
                        "status": 6,
                    },
                ),
            )
        )
    return records


def _extract_subscriptions(context: _WorkbookContext) -> list[LegacyRecord]:
    sheet_name = "账号订阅"
    formula_sheet = context.formulas[sheet_name]
    value_sheet = context.values[sheet_name]
    records: list[LegacyRecord] = []
    for row_number in _business_rows(
        formula_sheet, value_sheet, "AccountSubscriptionTable", (1, 6, 7)
    ):
        original_amount = _required_decimal(value_sheet.cell(row_number, 6).value)
        currency = _required_text(value_sheet.cell(row_number, 7).value).upper()
        fx_rate = _optional_decimal(value_sheet.cell(row_number, 8).value)
        if fx_rate is None:
            fx_rate = _lookup_rate(context.values["汇率参数"], currency, 5, 8)
        amount = _optional_decimal(value_sheet.cell(row_number, 9).value)
        if amount is None:
            amount = _round_money(original_amount * fx_rate)
        date_fields = _normalize_date(value_sheet.cell(row_number, 1).value)
        module = _optional_text(value_sheet.cell(row_number, 2).value)
        counterparty = _optional_text(value_sheet.cell(row_number, 3).value)
        expense_type = _optional_text(value_sheet.cell(row_number, 4).value)
        channel = _optional_text(value_sheet.cell(row_number, 5).value)
        notes = _optional_text(value_sheet.cell(row_number, 10).value)
        description = " / ".join(value for value in (module, counterparty) if value) or "账号订阅"
        business = {
            "record_type": "expense",
            "source_system": "subscription",
            "accounting_date": date_fields[0],
            "description": description,
            "original_amount": _decimal_string(original_amount),
            "currency": currency,
            "amount_cny": _decimal_string(amount),
        }
        source_fp, business_fp, fingerprint = _fingerprints(
            context=context,
            sheet=sheet_name,
            row=row_number,
            cell=None,
            business=business,
            trusted_external_id=None,
        )
        records.append(
            LegacyRecord(
                record_type="expense",
                source_system="subscription",
                category="subscription",
                accounting_date=date_fields[0],
                occurred_at=date_fields[1],
                date_precision=date_fields[2],
                service_period=date_fields[3],
                description=description,
                counterparty=counterparty,
                channel=channel,
                external_id=None,
                status=None,
                currency=currency,
                original_amount=_decimal_string(original_amount),
                fx_rate=_decimal_string(fx_rate),
                gross_amount_cny=None,
                fee_amount_cny=None,
                net_amount_cny=None,
                operating_amount_cny=_decimal_string(amount),
                cash_amount_cny=_decimal_string(amount),
                credited_value=None,
                recognition_method=None,
                notes=notes,
                fingerprint=fingerprint,
                business_fingerprint=business_fp,
                source_fingerprint=source_fp,
                evidence=_row_evidence(
                    context,
                    sheet_name,
                    row_number,
                    {
                        "occurred_at": 1,
                        "module": 2,
                        "counterparty": 3,
                        "expense_type": 4,
                        "channel": 5,
                        "original_amount": 6,
                        "currency": 7,
                        "fx_rate": 8,
                        "amount_cny": 9,
                        "notes": 10,
                    },
                ),
                metadata={"expense_type": expense_type, "module": module},
            )
        )
    return records


def _extract_other_expenses(context: _WorkbookContext) -> list[LegacyRecord]:
    sheet_name = "其他费用"
    formula_sheet = context.formulas[sheet_name]
    value_sheet = context.values[sheet_name]
    records: list[LegacyRecord] = []
    for row_number in _business_rows(formula_sheet, value_sheet, "OtherExpenseTable", (1, 2, 5)):
        original_amount = _required_decimal(value_sheet.cell(row_number, 5).value)
        currency = _required_text(value_sheet.cell(row_number, 4).value).upper()
        fx_rate = _optional_decimal(value_sheet.cell(row_number, 6).value)
        if fx_rate is None:
            fx_rate = _lookup_rate(context.values["汇率"], currency, 2, 4)
        amount = _optional_decimal(value_sheet.cell(row_number, 7).value)
        if amount is None:
            amount = _round_money(original_amount * fx_rate)
        date_fields = _normalize_date(value_sheet.cell(row_number, 1).value)
        description = _required_text(value_sheet.cell(row_number, 2).value)
        counterparty = _optional_text(value_sheet.cell(row_number, 3).value)
        notes = _optional_text(value_sheet.cell(row_number, 8).value)
        business = {
            "record_type": "expense",
            "source_system": "other_expense",
            "accounting_date": date_fields[0],
            "description": description,
            "original_amount": _decimal_string(original_amount),
            "currency": currency,
            "amount_cny": _decimal_string(amount),
        }
        source_fp, business_fp, fingerprint = _fingerprints(
            context=context,
            sheet=sheet_name,
            row=row_number,
            cell=None,
            business=business,
            trusted_external_id=None,
        )
        records.append(
            LegacyRecord(
                record_type="expense",
                source_system="other_expense",
                category="other_expense",
                accounting_date=date_fields[0],
                occurred_at=date_fields[1],
                date_precision=date_fields[2],
                service_period=date_fields[3],
                description=description,
                counterparty=counterparty,
                channel=None,
                external_id=None,
                status=None,
                currency=currency,
                original_amount=_decimal_string(original_amount),
                fx_rate=_decimal_string(fx_rate),
                gross_amount_cny=None,
                fee_amount_cny=None,
                net_amount_cny=None,
                operating_amount_cny=_decimal_string(amount),
                cash_amount_cny=_decimal_string(amount),
                credited_value=None,
                recognition_method=None,
                notes=notes,
                fingerprint=fingerprint,
                business_fingerprint=business_fp,
                source_fingerprint=source_fp,
                evidence=_row_evidence(
                    context,
                    sheet_name,
                    row_number,
                    {
                        "occurred_at": 1,
                        "description": 2,
                        "counterparty": 3,
                        "currency": 4,
                        "original_amount": 5,
                        "fx_rate": 6,
                        "amount_cny": 7,
                        "notes": 8,
                    },
                ),
            )
        )
    return records


def _extract_cloud_services(context: _WorkbookContext) -> list[LegacyRecord]:
    sheet_name = "云服务"
    formula_sheet = context.formulas[sheet_name]
    value_sheet = context.values[sheet_name]
    records: list[LegacyRecord] = []
    for row_number in _business_rows(formula_sheet, value_sheet, "CloudServiceTable", (1, 6, 8)):
        operating_original = _required_decimal(value_sheet.cell(row_number, 6).value)
        cash_original = _required_decimal(value_sheet.cell(row_number, 7).value)
        currency = _required_text(value_sheet.cell(row_number, 8).value).upper()
        fx_rate = _optional_decimal(value_sheet.cell(row_number, 9).value)
        if fx_rate is None:
            fx_rate = (
                Decimal("1") if currency == "CNY" else _required_decimal(value_sheet["B2"].value)
            )
        operating = _optional_decimal(value_sheet.cell(row_number, 10).value)
        cash = _optional_decimal(value_sheet.cell(row_number, 11).value)
        if operating is None:
            operating = (
                _round_money(operating_original * fx_rate)
                if currency == "USD"
                else operating_original
            )
        if cash is None:
            cash = _round_money(cash_original * fx_rate) if currency == "USD" else cash_original
        date_fields = _normalize_date(value_sheet.cell(row_number, 1).value)
        counterparty = _optional_text(value_sheet.cell(row_number, 2).value)
        description = _required_text(value_sheet.cell(row_number, 3).value)
        external_id = _optional_text(value_sheet.cell(row_number, 4).value)
        expense_type = _optional_text(value_sheet.cell(row_number, 5).value)
        status = _optional_text(value_sheet.cell(row_number, 12).value)
        business = {
            "record_type": "expense",
            "source_system": "cloud_service",
            "accounting_date": date_fields[0],
            "counterparty": counterparty,
            "description": description,
            "external_id": external_id,
            "operating": _decimal_string(operating),
            "cash": _decimal_string(cash),
        }
        source_fp, business_fp, fingerprint = _fingerprints(
            context=context,
            sheet=sheet_name,
            row=row_number,
            cell=None,
            business=business,
            trusted_external_id=None,
        )
        records.append(
            LegacyRecord(
                record_type="expense",
                source_system="cloud_service",
                category="cloud_service",
                accounting_date=date_fields[0],
                occurred_at=date_fields[1],
                date_precision=date_fields[2],
                service_period=date_fields[3],
                description=description,
                counterparty=counterparty,
                channel=None,
                external_id=external_id,
                status=status,
                currency=currency,
                original_amount=_decimal_string(operating_original),
                fx_rate=_decimal_string(fx_rate),
                gross_amount_cny=None,
                fee_amount_cny=None,
                net_amount_cny=None,
                operating_amount_cny=_decimal_string(operating),
                cash_amount_cny=_decimal_string(cash),
                credited_value=None,
                recognition_method=None,
                notes=None,
                fingerprint=fingerprint,
                business_fingerprint=business_fp,
                source_fingerprint=source_fp,
                evidence=_row_evidence(
                    context,
                    sheet_name,
                    row_number,
                    {
                        "period_or_date": 1,
                        "counterparty": 2,
                        "description": 3,
                        "external_id": 4,
                        "expense_type": 5,
                        "operating_original": 6,
                        "cash_original": 7,
                        "currency": 8,
                        "fx_rate": 9,
                        "operating_amount_cny": 10,
                        "cash_amount_cny": 11,
                        "status": 12,
                    },
                ),
                metadata={
                    "expense_type": expense_type,
                    "cash_original": _decimal_string(cash_original),
                    "unreconciled_difference_cny": _decimal_string(operating - cash),
                },
            )
        )
    return records


def _extract_upstream_recharges(context: _WorkbookContext) -> list[LegacyRecord]:
    sheet_name = "上游中转站充值"
    formula_sheet = context.formulas[sheet_name]
    value_sheet = context.values[sheet_name]
    table = formula_sheet.tables["UpstreamDailyTable"]
    min_column, min_row, max_column, max_row = range_boundaries(table.ref)
    provider_columns = range(min_column + 2, max_column)
    records: list[LegacyRecord] = []
    for row_number in range(min_row + 1, max_row + 1):
        raw_date = value_sheet.cell(row_number, min_column + 1).value
        if raw_date in (None, ""):
            continue
        date_fields = _normalize_date(raw_date)
        for column in provider_columns:
            amount = _optional_decimal(value_sheet.cell(row_number, column).value)
            if amount is None or amount == 0:
                continue
            provider = _required_text(value_sheet.cell(min_row, column).value)
            amount_cell = formula_sheet.cell(row_number, column).coordinate
            business = {
                "record_type": "upstream_recharge",
                "accounting_date": date_fields[0],
                "provider": provider,
                "amount": _decimal_string(amount),
            }
            source_fp, business_fp, fingerprint = _fingerprints(
                context=context,
                sheet=sheet_name,
                row=row_number,
                cell=amount_cell,
                business=business,
                trusted_external_id=None,
            )
            evidence = (
                _cell_evidence(context, sheet_name, min_row, column, "provider"),
                _cell_evidence(context, sheet_name, row_number, min_column + 1, "accounting_date"),
                _cell_evidence(context, sheet_name, row_number, column, "amount_cny"),
            )
            records.append(
                LegacyRecord(
                    record_type="upstream_recharge",
                    source_system="upstream",
                    category="upstream_recharge",
                    accounting_date=date_fields[0],
                    occurred_at=date_fields[1],
                    date_precision=date_fields[2],
                    service_period=date_fields[3],
                    description=f"上游充值 - {provider}",
                    counterparty=provider,
                    channel=None,
                    external_id=None,
                    status=None,
                    currency="CNY",
                    original_amount=_decimal_string(amount),
                    fx_rate="1",
                    gross_amount_cny=None,
                    fee_amount_cny=None,
                    net_amount_cny=None,
                    operating_amount_cny=_decimal_string(amount),
                    cash_amount_cny=_decimal_string(amount),
                    credited_value=_decimal_string(amount),
                    recognition_method="recharge_date",
                    notes=None,
                    fingerprint=fingerprint,
                    business_fingerprint=business_fp,
                    source_fingerprint=source_fp,
                    evidence=evidence,
                    metadata={"credit_currency": "UPSTREAM_CREDIT", "credit_fx_rate": "1"},
                )
            )
    return records


def _business_rows(
    formula_sheet: Worksheet,
    value_sheet: Worksheet,
    table_name: str,
    required_columns: tuple[int, ...],
) -> tuple[int, ...]:
    table = formula_sheet.tables[table_name]
    _min_column, min_row, _max_column, max_row = range_boundaries(table.ref)
    return tuple(
        row_number
        for row_number in range(min_row + 1, max_row + 1)
        if all(
            value_sheet.cell(row_number, column).value not in (None, "")
            for column in required_columns
        )
    )


def _row_evidence(
    context: _WorkbookContext,
    sheet_name: str,
    row_number: int,
    fields: dict[str, int],
) -> tuple[FieldEvidence, ...]:
    return tuple(
        _cell_evidence(context, sheet_name, row_number, column, field)
        for field, column in fields.items()
    )


def _cell_evidence(
    context: _WorkbookContext,
    sheet_name: str,
    row_number: int,
    column: int,
    field: str,
) -> FieldEvidence:
    formula_cell = context.formulas[sheet_name].cell(row_number, column)
    value_cell = context.values[sheet_name].cell(row_number, column)
    formula = str(formula_cell.value) if formula_cell.data_type == "f" else None
    value = value_cell.value if formula is not None else formula_cell.value
    return FieldEvidence(
        field=field,
        source=SourceLocation(
            file=context.source_label,
            sha256=context.sha256,
            sheet=sheet_name,
            row=row_number,
            column=column,
            cell=formula_cell.coordinate,
        ),
        value=normalize_scalar(value),
        formula=formula,
        comment=formula_cell.comment.text if formula_cell.comment else None,
    )


def _income_fx_rate(sheet: Worksheet, currency: str) -> Decimal:
    if currency == "CNY":
        return Decimal("1")
    if currency == "GBP":
        return _required_decimal(sheet["B3"].value)
    raise ValueError(f"Unsupported legacy income currency: {currency}")


def _lookup_rate(sheet: Worksheet, currency: str, start_row: int, end_row: int) -> Decimal:
    for row_number in range(start_row, end_row + 1):
        if _optional_text(sheet.cell(row_number, 1).value) == currency:
            return _required_decimal(sheet.cell(row_number, 2).value)
    raise ValueError(f"Missing exchange rate for {currency} in {sheet.title}")


def _normalize_date(
    value: Any,
) -> tuple[str, str | None, DatePrecision, str | None]:
    if isinstance(value, datetime):
        occurred_at = value.isoformat(sep=" ", timespec="seconds")
        return value.date().isoformat(), occurred_at, "second", None
    if isinstance(value, date):
        return value.isoformat(), None, "date", None

    text = _required_text(value).strip()
    month_match = _MONTH_PATTERN.fullmatch(text)
    if month_match:
        year, month = (int(part) for part in month_match.groups())
        last_day = calendar.monthrange(year, month)[1]
        return f"{year:04d}-{month:02d}-{last_day:02d}", None, "month", text
    datetime_match = _DATETIME_PATTERN.fullmatch(text)
    if datetime_match:
        day, clock = datetime_match.groups()
        if len(clock) == 5:
            clock = f"{clock}:00"
        return day, f"{day} {clock}", "second", None
    if _DATE_PATTERN.fullmatch(text):
        return text, None, "date", None
    raise ValueError(f"Unsupported legacy date value: {value!r}")


def _fingerprints(
    *,
    context: _WorkbookContext,
    sheet: str,
    row: int,
    cell: str | None,
    business: dict[str, Any],
    trusted_external_id: str | None,
) -> tuple[str, str, str]:
    source_fingerprint = _digest(
        {
            "file_sha256": context.sha256,
            "sheet": sheet,
            "row": row,
            "cell": cell,
        }
    )
    business_fingerprint = _digest(business)
    if trusted_external_id:
        fingerprint = _digest(
            {
                "record_type": business["record_type"],
                "source_system": business["source_system"],
                "external_id": trusted_external_id,
            }
        )
    else:
        fingerprint = _digest(
            {"source_fingerprint": source_fingerprint, "business_fingerprint": business_fingerprint}
        )
    return source_fingerprint, business_fingerprint, fingerprint


def _digest(value: dict[str, Any]) -> str:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _required_decimal(value: Any) -> Decimal:
    number = _optional_decimal(value)
    if number is None:
        raise ValueError("Required decimal value is empty")
    return number


def _optional_decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        raise ValueError("Boolean is not a decimal amount")
    return Decimal(str(value))


def _round_money(value: Decimal) -> Decimal:
    return value.quantize(_CENT, rounding=ROUND_HALF_UP)


def _decimal_string(value: Decimal) -> str:
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _required_text(value: Any) -> str:
    text = _optional_text(value)
    if text is None:
        raise ValueError("Required text value is empty")
    return text


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
