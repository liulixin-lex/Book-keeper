from __future__ import annotations

import csv
import hashlib
import mimetypes
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any
from uuid import uuid4

import openpyxl  # type: ignore[import-untyped]
import pypdfium2  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from fastapi import UploadFile
from pypdf import PdfReader
from rapidocr import RapidOCR
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import Settings
from app.enums import AIWorkflowStatus
from app.models import AIWorkflow, AIWorkflowEvent, AIWorkflowInput, DocumentExtraction, User

PARSER_VERSION = "document-extractor-1"
SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".xls", ".xlsx", ".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}


class DocumentExtractionError(ValueError):
    def __init__(self, message: str, *, status_code: int = 422) -> None:
        super().__init__(message)
        self.status_code = status_code


def _safe_upload_path(settings: Settings, workflow_id: str, filename: str) -> Path:
    root = settings.ai_upload_dir.resolve()
    workflow_dir = (root / workflow_id).resolve()
    if root not in workflow_dir.parents:
        raise DocumentExtractionError("上传目录无效")
    workflow_dir.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix.lower()
    return workflow_dir / f"{uuid4().hex}{suffix}"


def _delete_temporary(path: str | None, settings: Settings) -> None:
    if not path:
        return
    root = settings.ai_upload_dir.resolve()
    target = Path(path).resolve()
    if root not in target.parents:
        return
    target.unlink(missing_ok=True)
    parent = target.parent
    if parent != root:
        try:
            parent.rmdir()
        except OSError:
            pass


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("文本编码不受支持")


def _extract_text(path: Path) -> tuple[str, dict[str, Any], str]:
    content = _decode_text(path.read_bytes())
    locators = [
        {"line": index, "text": line}
        for index, line in enumerate(content.splitlines(), start=1)
        if line.strip()
    ]
    return content, {"locators": locators}, "text"


def _extract_csv(path: Path) -> tuple[str, dict[str, Any], str]:
    content = _decode_text(path.read_bytes())
    rows = list(csv.reader(content.splitlines()))
    lines: list[str] = []
    cells: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=1):
        values: list[str] = []
        for column_index, value in enumerate(row, start=1):
            normalized = value.strip()
            values.append(normalized)
            if normalized:
                cells.append({"row": row_index, "column": column_index, "value": normalized})
        lines.append("\t".join(values))
    return "\n".join(lines), {"cells": cells}, "csv"


def _extract_xlsx(path: Path) -> tuple[str, dict[str, Any], str]:
    values_book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    formulas_book = openpyxl.load_workbook(path, read_only=False, data_only=False)
    lines: list[str] = []
    cells: list[dict[str, Any]] = []
    try:
        for sheet_name in values_book.sheetnames:
            values_sheet = values_book[sheet_name]
            formulas_sheet = formulas_book[sheet_name]
            lines.append(f"[sheet:{sheet_name}]")
            for row in values_sheet.iter_rows():
                row_values: list[str] = []
                for cell in row:
                    formula_cell = formulas_sheet[cell.coordinate]
                    formula = (
                        str(formula_cell.value)
                        if formula_cell.data_type == "f" and formula_cell.value is not None
                        else None
                    )
                    value = cell.value
                    normalized = "" if value is None else str(value)
                    row_values.append(normalized)
                    if normalized or formula:
                        cells.append(
                            {
                                "sheet": sheet_name,
                                "cell": cell.coordinate,
                                "value": value,
                                "formula": formula,
                                "formula_result_missing": bool(formula and value is None),
                            }
                        )
                if any(row_values):
                    lines.append("\t".join(row_values))
    finally:
        values_book.close()
        formulas_book.close()
    return "\n".join(lines), {"cells": cells, "formulas_executed": False}, "xlsx"


def _extract_xls(path: Path) -> tuple[str, dict[str, Any], str]:
    book = xlrd.open_workbook(path, on_demand=True)
    lines: list[str] = []
    cells: list[dict[str, Any]] = []
    try:
        for sheet in book.sheets():
            lines.append(f"[sheet:{sheet.name}]")
            for row_index in range(sheet.nrows):
                values: list[str] = []
                for column_index in range(sheet.ncols):
                    value = sheet.cell_value(row_index, column_index)
                    normalized = "" if value in {None, ""} else str(value)
                    values.append(normalized)
                    if normalized:
                        cells.append(
                            {
                                "sheet": sheet.name,
                                "row": row_index + 1,
                                "column": column_index + 1,
                                "value": value,
                            }
                        )
                if any(values):
                    lines.append("\t".join(values))
    finally:
        book.release_resources()
    return "\n".join(lines), {"cells": cells, "formulas_executed": False}, "xls"


def _ocr_image(image: Any) -> tuple[str, list[dict[str, Any]]]:
    result = RapidOCR()(image)
    texts = getattr(result, "txts", None)
    boxes = getattr(result, "boxes", None)
    scores = getattr(result, "scores", None)
    if texts is None and isinstance(result, tuple) and result:
        rows = result[0] or []
        texts = [str(row[1]) for row in rows]
        boxes = [row[0] for row in rows]
        scores = [row[2] for row in rows]
    texts = list(texts or [])
    boxes = list(boxes or [None] * len(texts))
    scores = list(scores or [None] * len(texts))
    locators = [
        {"text": str(text), "bbox": box, "score": score}
        for text, box, score in zip(texts, boxes, scores, strict=False)
    ]
    return "\n".join(str(value) for value in texts), locators


def _extract_pdf(path: Path) -> tuple[str, dict[str, Any], str]:
    reader = PdfReader(path)
    pages: list[str] = []
    locators: list[dict[str, Any]] = []
    ocr_pages: list[int] = []
    pdfium = None
    try:
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            page_locators: list[dict[str, Any]] = []
            if len(text) < 20:
                if pdfium is None:
                    pdfium = pypdfium2.PdfDocument(str(path))
                image = pdfium[index - 1].render(scale=2).to_pil()
                text, page_locators = _ocr_image(image)
                ocr_pages.append(index)
            pages.append(f"[page:{index}]\n{text}")
            locators.append({"page": index, "text": text, "ocr": index in ocr_pages})
            for item in page_locators:
                locators.append({"page": index, **item, "ocr": True})
    finally:
        if pdfium is not None:
            pdfium.close()
    return "\n\n".join(pages), {"locators": locators, "ocr_pages": ocr_pages}, "pdf"


def _extract_image(path: Path) -> tuple[str, dict[str, Any], str]:
    text, locators = _ocr_image(str(path))
    return text, {"locators": locators}, "image_ocr"


def extract_document(path: Path) -> tuple[str, dict[str, Any], str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _extract_text(path)
    if suffix == ".csv":
        return _extract_csv(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    if suffix == ".xls":
        return _extract_xls(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in IMAGE_SUFFIXES:
        return _extract_image(path)
    raise DocumentExtractionError("不支持的附件类型")


def _next_event_sequence(session: Session, workflow_id: str) -> int:
    current = session.scalar(
        select(func.coalesce(func.max(AIWorkflowEvent.sequence), 0)).where(
            AIWorkflowEvent.workflow_id == workflow_id
        )
    )
    return int(current or 0) + 1


async def add_and_extract_upload(
    session: Session,
    user: User,
    workflow: AIWorkflow,
    upload: UploadFile,
    settings: Settings,
) -> AIWorkflowInput:
    if workflow.workspace_id != user.workspace_id or workflow.created_by_id != user.id:
        raise DocumentExtractionError("工作流不存在", status_code=404)
    if AIWorkflowStatus(workflow.status) not in {
        AIWorkflowStatus.COLLECTING,
        AIWorkflowStatus.NEEDS_INPUT,
        AIWorkflowStatus.DRAFT,
    }:
        raise DocumentExtractionError("当前状态不能上传附件", status_code=409)
    filename = Path(upload.filename or "attachment").name
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES | IMAGE_SUFFIXES:
        raise DocumentExtractionError("不支持的附件类型")
    count, total_size = session.execute(
        select(
            func.count(AIWorkflowInput.id), func.coalesce(func.sum(AIWorkflowInput.size_bytes), 0)
        ).where(AIWorkflowInput.workflow_id == workflow.id, AIWorkflowInput.kind == "file")
    ).one()
    if int(count) >= settings.ai_max_files_per_workflow:
        raise DocumentExtractionError("附件数量超过当前工作流上限", status_code=413)
    destination = _safe_upload_path(settings, workflow.id, filename)
    digest = hashlib.sha256()
    size = 0
    try:
        with destination.open("xb") as stream:
            while chunk := await upload.read(1024 * 1024):
                size += len(chunk)
                if size > settings.ai_max_file_size_bytes:
                    raise DocumentExtractionError("单个附件超过大小上限", status_code=413)
                if int(total_size) + size > settings.ai_max_total_file_size_bytes:
                    raise DocumentExtractionError("附件总大小超过当前工作流上限", status_code=413)
                digest.update(chunk)
                stream.write(chunk)
    except Exception:
        _delete_temporary(str(destination), settings)
        raise
    sequence = (
        session.scalar(
            select(func.coalesce(func.max(AIWorkflowInput.sequence), 0)).where(
                AIWorkflowInput.workflow_id == workflow.id
            )
        )
        or 0
    ) + 1
    item = AIWorkflowInput(
        workflow_id=workflow.id,
        sequence=sequence,
        kind="file",
        filename=filename,
        mime_type=upload.content_type
        or mimetypes.guess_type(filename)[0]
        or "application/octet-stream",
        sha256=digest.hexdigest(),
        size_bytes=size,
        temporary_path=str(destination),
        status="extracting",
    )
    session.add(item)
    workflow.status = AIWorkflowStatus.EXTRACTING
    workflow.version += 1
    workflow.canonical_hash = None
    session.flush()
    try:
        content, structure, parser = extract_document(destination)
        extraction = DocumentExtraction(
            workflow_input_id=item.id,
            parser=parser,
            parser_version=PARSER_VERSION,
            content_text=content,
            structure_json=structure,
            status="completed",
        )
        session.add(extraction)
        item.status = "extracted"
        item.temporary_path = None
        workflow.status = AIWorkflowStatus.COLLECTING
        session.add(
            AIWorkflowEvent(
                workflow_id=workflow.id,
                sequence=_next_event_sequence(session, workflow.id),
                event_type="document.extracted",
                payload_json={
                    "input_id": item.id,
                    "filename": filename,
                    "sha256": item.sha256,
                    "parser": parser,
                    "parser_version": PARSER_VERSION,
                },
            )
        )
        _delete_temporary(str(destination), settings)
    except Exception as exc:
        item.status = "failed"
        workflow.status = AIWorkflowStatus.COLLECTING
        session.add(
            DocumentExtraction(
                workflow_input_id=item.id,
                parser="unknown",
                parser_version=PARSER_VERSION,
                content_text="",
                structure_json={},
                status="failed",
                error=str(exc)[:2000],
            )
        )
        session.add(
            AIWorkflowEvent(
                workflow_id=workflow.id,
                sequence=_next_event_sequence(session, workflow.id),
                event_type="document.failed",
                payload_json={"input_id": item.id, "filename": filename, "error": str(exc)[:500]},
            )
        )
    session.commit()
    session.refresh(item)
    return item


def cleanup_failed_uploads(session: Session, settings: Settings) -> int:
    cutoff = datetime.now(UTC) - timedelta(hours=settings.ai_failed_file_ttl_hours)
    items = session.scalars(
        select(AIWorkflowInput).where(
            AIWorkflowInput.status == "failed",
            AIWorkflowInput.temporary_path.is_not(None),
            AIWorkflowInput.created_at < cutoff,
        )
    ).all()
    for item in items:
        _delete_temporary(item.temporary_path, settings)
        item.temporary_path = None
        item.status = "expired"
    if items:
        session.commit()
    return len(items)
