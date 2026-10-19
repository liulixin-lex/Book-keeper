"""Parser framework: pure functions from file bytes (or JSON records) to normalised rows (docs/06 §2)."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Protocol
from zoneinfo import ZoneInfo


class ParseError(ValueError):
    """The file as a whole cannot be read (wrong format, missing header)."""


@dataclass(frozen=True)
class ParsedRecord:
    row_no: int
    occurred_at: datetime  # timezone-aware
    direction: str  # in | out | neutral
    amount: Decimal  # positive, major units
    commodity: str
    external_id: str | None = None
    fee_amount: Decimal | None = None
    balance_after: Decimal | None = None
    counterparty_text: str | None = None
    description: str | None = None
    order_no: str | None = None
    reference: str | None = None
    status_text: str | None = None
    extracted: Mapping[str, Any] = field(default_factory=dict)
    raw: Mapping[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class RowError:
    row_no: int
    message: str
    raw: Mapping[str, Any] = field(default_factory=dict)


@dataclass
class FileSummary:
    period_from: date | None = None
    period_to: date | None = None
    declared: dict[str, Decimal] = field(default_factory=dict)  # in_count, in_amount, out_count, out_amount
    opening_balance: Decimal | None = None
    closing_balance: Decimal | None = None
    account_identifier: str | None = None


@dataclass
class ParseResult:
    records: list[ParsedRecord] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    skipped: list[RowError] = field(default_factory=list)  # rows deliberately left out (closed / refunded …)
    summary: FileSummary = field(default_factory=FileSummary)


class Parser(Protocol):
    code: str
    version: str

    def sniff(self, filename: str, head: bytes) -> float: ...

    def parse(self, data: bytes, filename: str, config: Mapping[str, Any], tz: ZoneInfo) -> ParseResult: ...


# ------------------------------------------------------------------ reading helpers


def decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文件编码（支持 UTF-8 与 GBK）。")


def clean(value: Any) -> str:
    """Strip whitespace, tabs and the leading apostrophes / tabs wallets add to keep numbers as text."""
    if value is None:
        return ""
    text = str(value).replace(" ", " ").replace("​", "").strip()
    return text.strip("\t '`").strip()


def read_rows(data: bytes, filename: str) -> list[list[str]]:
    """CSV / TSV / XLSX / JSON (list of objects) → rows of cleaned strings (header included)."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    if name.endswith(".json"):
        items = json.loads(decode(data))
        if isinstance(items, dict):
            items = items.get("records") or items.get("data") or []
        if not isinstance(items, list) or not all(isinstance(i, dict) for i in items):
            raise ParseError("JSON 必须是对象数组，或包含 records 数组的对象。")
        keys: list[str] = []
        for item in items:
            keys.extend(k for k in item if k not in keys)
        return [keys] + [[_cell(item.get(k)) for k in keys] for item in items]
    text = decode(data)
    sample = text[:4096]
    delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    return [[clean(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _cell(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)).normalize(), "f")
    return clean(value)


def find_header(rows: Sequence[Sequence[str]], required: Sequence[str], limit: int = 40) -> int:
    """Index of the first row containing every required column name."""
    for index, row in enumerate(rows[:limit]):
        cells = {c.strip() for c in row}
        if all(any(r == c or r in c for c in cells) for r in required):
            return index
    raise ParseError(f"未找到表头（需要列：{'、'.join(required)}）。")


def column_index(header: Sequence[str], *names: str) -> int | None:
    for name in names:
        for i, cell in enumerate(header):
            if cell.strip() == name:
                return i
    for name in names:
        for i, cell in enumerate(header):
            if name in cell:
                return i
    return None


AMOUNT_RE = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?")


def parse_amount(value: str) -> Decimal | None:
    """'¥51.50' / '-1,234.5' / '51.5元' → Decimal; '' → None."""
    text = clean(value).replace("￥", "").replace("¥", "").replace("$", "").replace("£", "")
    if not text or text in ("/", "-", "--"):
        return None
    match = AMOUNT_RE.search(text)
    if not match:
        raise ValueError(f"金额无法识别：{value}")
    try:
        return Decimal(match.group(0).replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"金额无法识别：{value}") from exc


DATETIME_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M:%S",
    "%Y/%m/%d %H:%M",
    "%Y-%m-%dT%H:%M:%S",
    "%Y年%m月%d日 %H:%M:%S",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y年%m月%d日",
)
EXCEL_EPOCH = datetime(1899, 12, 30)  # noqa: DTZ001 - naive by definition; tz applied after


def parse_datetime(value: str, tz: ZoneInfo, fmt: str | None = None) -> datetime:
    """Local time in ``tz`` unless the value carries an offset; Excel serial numbers are accepted."""
    text = clean(value)
    if not text:
        raise ValueError("时间为空")
    if re.fullmatch(r"\d{5}(\.\d+)?", text):
        return (EXCEL_EPOCH + timedelta(days=float(text))).replace(tzinfo=tz)
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=tz)
    except ValueError:
        pass
    for candidate in ((fmt,) if fmt else ()) + DATETIME_FORMATS:
        try:
            return datetime.strptime(text, candidate).replace(tzinfo=tz)
        except ValueError:
            continue
    raise ValueError(f"时间无法识别：{value}")


def mask_secret(value: str) -> dict[str, str]:
    """Redemption codes / card secrets: keep only a SHA-256 digest and the last four characters."""
    text = clean(value)
    return {"sha256": hashlib.sha256(text.encode()).hexdigest(), "last4": text[-4:]}


FACE_PATTERNS = (
    r"余额\s*\+\s*(?P<face>\d+(?:\.\d+)?)",
    r"(?P<face>\d+(?:\.\d+)?)\s*[RrＲ](?![a-zA-Z])",
    r"(?P<face>\d+(?:\.\d+)?)\s*(?:刀|美元|美金|USD|\$)",
    r"\$\s*(?P<face>\d+(?:\.\d+)?)",
    r"(?P<face>\d+(?:\.\d+)?)\s*额度",
)


def extract_face(text: str | None, patterns: Sequence[str] = FACE_PATTERNS) -> Decimal | None:
    if not text:
        return None
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return Decimal(match.group("face") if "face" in match.groupdict() else match.group(1))
    return None
