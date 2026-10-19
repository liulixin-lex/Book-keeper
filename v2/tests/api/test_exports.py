"""Exports (CSV, XLSX, Beancount, JSON), MCP export, monthly archive."""

from __future__ import annotations

import csv
import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

from beancount import loader
from beancount.core import data as bean
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bookkeeper.db.models import Book
from bookkeeper.services import exports
from tests.api.conftest import login, write_headers
from tests.api.test_business import biz, money


def setup_october(client, book: str, csrf: str) -> None:
    capital = client.post(
        f"/api/v1/books/{book}/transactions",
        headers=write_headers(csrf),
        json={
            "type": "manual_journal",
            "effective_date": "2026-10-01",
            "description": "调整 · 股东投入",
            "lines": [
                {"account": "1012.01", "side": "debit", "amount": money("1000.00")},
                {"account": "3001", "side": "credit", "amount": money("1000.00")},
            ],
        },
    )
    assert capital.status_code == 201, capital.text
    biz(
        client,
        book,
        csrf,
        "sale",
        {"gross": money("20.00", "GBP"), "fee": money("0.85", "GBP"), "channel": "stripe", "fx_rate": "9.5"},
        effective_date="2026-10-10",
    )
    biz(client, book, csrf, "sale", {"gross": money("100.00"), "channel": "liandong"}, effective_date="2026-10-11")
    biz(client, book, csrf, "expense", {"amount": money("30.00"), "category": "5401.04"}, effective_date="2026-10-12")


def rich_book(client, book: str, csrf: str) -> None:
    """Every kind of entry: foreign sale, conversion with fx realization, a deletion and an edit (re-posted)."""
    setup_october(client, book, csrf)
    client.post(
        f"/api/v1/books/{book}/fx-rates",
        headers=write_headers(csrf),
        json={"base": "GBP", "quote": "CNY", "rate_date": "2026-10-20", "rate": "9.6"},
    )
    biz(
        client,
        book,
        csrf,
        "transfer",
        {
            "from_account": "1012.01",
            "to_account": "1012.01",
            "amount_out": money("10.00", "GBP"),
            "amount_in": money("95.00"),
        },
        effective_date="2026-10-20",
    )
    wrong = biz(
        client, book, csrf, "expense", {"amount": money("12.00"), "category": "5602.99"}, effective_date="2026-10-21"
    )
    deleted = client.delete(f"/api/v1/books/{book}/transactions/{wrong['number']}", headers=write_headers(csrf))
    assert deleted.status_code == 200, deleted.text
    sale = client.get(f"/api/v1/books/{book}/transactions", params={"channel": "liandong"}).json()["data"][0]
    edited = client.patch(
        f"/api/v1/books/{book}/transactions/{sale['number']}",
        headers=write_headers(csrf),
        json={"inputs": {"gross": money("110.00")}},
    )
    assert edited.status_code == 200, edited.text


def test_exports_are_valid_and_consistent(client, world) -> None:
    csrf = login(client, world.admin_email)
    book = world.book_id
    rich_book(client, book, csrf)

    response = client.get(f"/api/v1/books/{book}/exports/beancount")
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"].endswith(".beancount")
    text = response.text
    entries, errors, _options = loader.load_string(text)
    assert errors == [], errors[:3]
    txns = [e for e in entries if isinstance(e, bean.Transaction)]
    lines = int(response.headers["x-export-lines"])
    assert len(txns) == 9 and sum(len(t.postings) for t in txns) == lines
    assert any(p.account.endswith("-汇兑损益") for t in txns for p in t.postings)
    assert any(p.price is not None and p.units.currency == "GBP" for t in txns for p in t.postings)

    rows = list(csv.reader(io.StringIO(client.get(f"/api/v1/books/{book}/exports/csv").content.decode("utf-8-sig"))))
    header, body = rows[0], rows[1:]
    assert header[:3] == ["凭证号", "业务日期", "分录类型"] and len(body) == lines
    base_debit = sum(Decimal(r[header.index("本位币借方")] or 0) for r in body)
    base_credit = sum(Decimal(r[header.index("本位币贷方")] or 0) for r in body)
    assert base_debit == base_credit > 0

    dump = json.loads(client.get(f"/api/v1/books/{book}/exports/json", params={"date_from": "2026-10-01"}).content)
    assert dump["book"]["base_commodity"] == "CNY" and len(dump["entries"]) == 9
    assert all(sum(Decimal(li["base_amount"]) for li in e["lines"]) == 0 for e in dump["entries"])

    workbook = load_workbook(io.BytesIO(client.get(f"/api/v1/books/{book}/exports/xlsx").content), read_only=True)
    assert workbook.sheetnames == ["分录", "科目"]
    assert sum(1 for _ in workbook["分录"].iter_rows()) == lines + 1

    narrow = client.get(
        f"/api/v1/books/{book}/exports/csv", params={"date_from": "2026-10-12", "date_to": "2026-10-12"}
    )
    assert narrow.headers["x-export-lines"] == "2"  # expense without fee: cost + pay
    bad = client.get(f"/api/v1/books/{book}/exports/csv", params={"date_from": "2026-10-12", "date_to": "2026-10-01"})
    assert bad.status_code == 422
    assert client.get(f"/api/v1/books/{book}/exports/pdf").status_code == 422


def test_mcp_export_and_download(client, world) -> None:
    from tests.api.test_mcp import mcp_for

    csrf = login(client, world.admin_email)
    setup_october(client, world.book_id, csrf)
    reader = mcp_for(client, world, csrf, preset="reader")
    assert "export_ledger" in reader.tools()
    result = reader.call("export_ledger", format="beancount", date_from="2026-10-01", date_to="2026-10-31")
    assert result["ok"], result
    auth = {"Authorization": reader.headers["Authorization"]}
    assert client.get(result["data"]["download"], headers=auth).status_code == 403  # reader preset lacks files:read
    download = client.get(result["data"]["export_url"], headers=auth)
    assert download.status_code == 200
    _entries, errors, _ = loader.load_string(download.text)
    assert errors == []
    again = reader.call("export_ledger", format="beancount", date_from="2026-10-01", date_to="2026-10-31")
    assert again["data"]["file"]["id"] == result["data"]["file"]["id"]  # same content, same file


def test_monthly_archive_rewrites_changed_months(client, world, app, tmp_path: Path) -> None:
    csrf = login(client, world.admin_email)
    setup_october(client, world.book_id, csrf)
    with app.state.sessionmaker() as session, session.begin():
        book = session.get(Book, world.book_id)
        assert exports.archive_months(session, book, tmp_path, date(2026, 11, 5)) == ["2026-10"]
        assert exports.archive_months(session, book, tmp_path, date(2026, 11, 6)) == []
        assert exports.archive_months(session, book, tmp_path, date(2026, 10, 20)) == []  # current month waits
    folder = tmp_path / world.book_id / "2026-10"
    manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["entries"] == 4 and set(manifest["sha256"]) == {"ledger.beancount", "journal.csv"}
    assert loader.load_string((folder / "ledger.beancount").read_text(encoding="utf-8"))[1] == []

    biz(
        client,
        world.book_id,
        csrf,
        "expense",
        {"amount": money("3.00"), "category": "5602.99"},
        effective_date="2026-10-30",
    )
    with app.state.sessionmaker() as session, session.begin():
        book = session.get(Book, world.book_id)
        assert exports.archive_months(session, book, tmp_path, date(2026, 11, 7)) == ["2026-10"]
